from __future__ import annotations

import csv
import io

from ai_test_asset_center.enterprise_knowledge_center import _parse_source


HEADERS = ["Table", "Field", "Type", "Description", "Required"]
ROWS = [
    ["orders", "order_id", "string", "order identifier", "yes"],
    ["orders", "amount", "decimal", "payable amount", "yes"],
    ["orders", "tenant_id", "string", "tenant scope", "no"],
]


def _markdown_bytes() -> bytes:
    lines = [
        "| " + " | ".join(HEADERS) + " |",
        "| " + " | ".join("---" for _ in HEADERS) + " |",
        *["| " + " | ".join(row) + " |" for row in ROWS],
    ]
    return "\n".join(lines).encode("utf-8")


def _csv_bytes() -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer)
    writer.writerow(HEADERS)
    writer.writerows(ROWS)
    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Field Dictionary"
    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    table = document.add_table(rows=1, cols=len(HEADERS))
    for index, value in enumerate(HEADERS):
        table.rows[0].cells[index].text = value
    for values in ROWS:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _semantic_shape(data: bytes, filename: str, source_id: str) -> tuple[set[str], set[str]]:
    parsed = _parse_source(
        data,
        filename,
        "db_field_dictionary",
        source_id,
    )
    return (
        {
            str(row.get("name") or "")
            for row in parsed.get("tables") or []
            if str(row.get("name") or "")
        },
        {
            str(row.get("field") or "")
            for row in parsed.get("field_dictionary") or []
            if str(row.get("field") or "")
        },
    )


def test_same_field_dictionary_has_equivalent_semantics_across_four_formats() -> None:
    shapes = {
        "md": _semantic_shape(_markdown_bytes(), "fields.md", "source-md"),
        "csv": _semantic_shape(_csv_bytes(), "fields.csv", "source-csv"),
        "xlsx": _semantic_shape(_xlsx_bytes(), "fields.xlsx", "source-xlsx"),
        "docx": _semantic_shape(_docx_bytes(), "fields.docx", "source-docx"),
    }

    assert shapes == {
        format_name: (
            {"orders"},
            {"order_id", "amount", "tenant_id"},
        )
        for format_name in shapes
    }
